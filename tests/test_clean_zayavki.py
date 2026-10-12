from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import clean_zayavki as cz
import make_test_file as gen

TODAY = date(2026, 10, 11)
ROOT = Path(__file__).resolve().parent.parent
ORIGINAL = ROOT / "zayavki.xlsx"


# ---------------------------------------------------------------- телефон

@pytest.mark.parametrize("raw, expected", [
    # все форматы из исходного файла
    ("8 (916) 245-17-83", "+79162451783"),
    ("+7 916 245 17 83", "+79162451783"),
    (89032218840, "+79032218840"),                   # число в Excel
    ("+7 (903) 221-88-40", "+79032218840"),
    ("7 925 600 31 12", "+79256003112"),             # 7 без плюса
    (79256003112, "+79256003112"),
    ("+7 985 112-90-47", "+79851129047"),
    ("тел.: 8 917 400 52 19 (после 18:00)", "+79174005219"),  # номер внутри текста
    ("+79174005219", "+79174005219"),
    ("8-916-555-66-77", "+79165556677"),
    # дополнительные случаи
    ("9161234567", "+79161234567"),                  # 10 цифр без кода
    (9161234567, "+79161234567"),
    ("(812) 345-67-89", "+78123456789"),             # СПб без 8
    ("8 (495) 123-45-67", "+74951234567"),
    (89169990011.0, "+79169990011"),                 # float из Excel
    ("+7 916 111-22-33 доб. 45", "+79161112233"),    # добавочный отбрасываем
    ("моб. 8 916 888 99 00, раб. 8 495 765 43 21", "+79168889900"),
    ("8.916.123.45.67", "+79161234567"),
    ("+7 916 123 45 67", "+79161234567"),  # неразрывные пробелы
    ("+7 727 123 45 67", "+77271234567"),            # Казахстан — тоже +7
])
def test_phone_ok(raw, expected):
    res = cz.normalize_phone(raw)
    assert res.value == expected and res.problem is None


@pytest.mark.parametrize("raw, reason", [
    (None, "пусто"),
    ("", "пусто"),
    ("   ", "пусто"),
    ("не знаю", "не номер"),
    ("abc", "не номер"),
    ("8 (800) 707-46-19", "8-80X"),                  # бесплатная линия организации
    ("88007074619", "8-80X"),
    ("+7 800 707-46-19", "8-80X"),
    ("8 804 333 22 11", "8-80X"),
    ("8 916 333 44 5", "не хватает цифры"),          # 10 цифр, начинается с 89
    ("7 925 600 31 1", "не хватает цифры"),
    ("916 123 45", "не хватает цифр: 8"),
    ("+7 916 123 45 6", "не хватает цифр: 10"),
    ("+7 916 444 55 66 7", "лишние цифры: 12"),
    ("8 916 123 45 67 8", "лишние цифры: 12"),
    ("+44 20 7946 0958", "иностранный номер"),
    ("+7 000 000 00 00", "несуществующий код"),
    ("89999999999", "заглушка"),
    ("+7 (812) 3O1-22-33", "не хватает цифр"),        # буква O вместо нуля
    (True, "не номер"),
])
def test_phone_bad(raw, reason):
    res = cz.normalize_phone(raw)
    assert res.value is None
    assert reason in res.problem


def test_phone_notes():
    assert "добавочный" in cz.normalize_phone("+7 916 111-22-33 доб. 45").note
    assert "несколько номеров" in cz.normalize_phone("8 916 888 99 00, 8 495 765 43 21").note
    assert cz.normalize_phone("+7 916 111-22-33").note is None


# ---------------------------------------------------------------- дата

@pytest.mark.parametrize("raw, expected", [
    ("24.09.2026", date(2026, 9, 24)),               # 03.09 неоднозначна, её решает resolve_*
    ("13.09.26", date(2026, 9, 13)),                 # двузначный год -> 20YY
    ("2026-09-08", date(2026, 9, 8)),
    ("2026/09/26", date(2026, 9, 26)),
    ("2026.08.12", date(2026, 8, 12)),
    ("17/09/2026", date(2026, 9, 17)),               # однозначно ДД/ММ
    ("09/16/2026", date(2026, 9, 16)),               # однозначно ММ/ДД
    ("09/09/2026", date(2026, 9, 9)),                # день = месяц
    ("6 сентября 2026", date(2026, 9, 6)),
    ("23 сентября 2026", date(2026, 9, 23)),
    ("5 авг. 2026 г.", date(2026, 8, 5)),
    ("1 сент 2026г.", date(2026, 9, 1)),
    ("1 мая 2026 года", date(2026, 5, 1)),
    ("Aug 10, 2026", date(2026, 8, 10)),
    ("10 September 2026", date(2026, 9, 10)),
    ("24.09.2026 14:35", date(2026, 9, 24)),         # время отбрасываем
    ("2026-09-08T14:30:00", date(2026, 9, 8)),
    (datetime(2026, 9, 6), date(2026, 9, 6)),        # настоящая дата Excel
    (datetime(2026, 8, 20, 15, 45), date(2026, 8, 20)),
    (date(2026, 9, 1), date(2026, 9, 1)),
    (46268, date(2026, 9, 3)),                       # Excel-дата числом
    (46268.75, date(2026, 9, 3)),
    ("46268", date(2026, 9, 3)),
])
def test_date_ok(raw, expected):
    res = cz.parse_date(raw, TODAY)
    assert res.value == expected and res.problem is None and res.ambiguous is None


@pytest.mark.parametrize("raw, reason", [
    (None, "пусто"),
    ("  ", "пусто"),
    ("хх.09.2026", "не распознана"),
    ("12 сентября", "нет года"),
    ("19.09.2062", "в будущем"),
    ("31.02.2026", "несуществующая"),
    ("01.01.1990", "старая"),
    ("сентябрь 2026", "не распознана"),
    ("2026-13-01", "несуществующая"),
    (123, "не распознана"),                          # число, но не похоже на Excel-дату
])
def test_date_bad(raw, reason):
    res = cz.parse_date(raw, TODAY)
    assert res.value is None
    assert reason in res.problem


def test_date_ambiguous_returns_both_variants():
    res = cz.parse_date("09/03/2026", TODAY)
    assert res.value is None and res.problem is None
    assert res.ambiguous == (date(2026, 3, 9), date(2026, 9, 3))  # (ДД/ММ, ММ/ДД)


def _recs(rows):
    return [cz.build_record(i, "Имя Фамилия", f"+7 916 000 00 {i:02d}", d, s, TODAY)
            for i, (d, s) in enumerate(rows, start=2)]


def test_ambiguous_resolved_by_source_evidence():
    recs = _recs([("09/16/2026", "сайт"), ("09/03/2026", "сайт"),
                  ("17/09/2026", "telegram"), ("07/09/2026", "telegram")])
    cz.resolve_ambiguous_dates(recs, TODAY)
    assert recs[1].date == date(2026, 9, 3)   # сайт пишет ММ/ДД
    assert recs[3].date == date(2026, 9, 7)   # telegram пишет ДД/ММ
    assert "сайт" in recs[1].date_res.note


def test_ambiguous_mixed_source_falls_back_to_proximity():
    # у источника есть и ДД/ММ, и ММ/ДД -> подсказки нет, берём ближайшее к медиане файла
    recs = _recs([("09/16/2026", "x"), ("17/09/2026", "x"), ("2026-09-20", "y"), ("09/04/2026", "x")])
    cz.resolve_ambiguous_dates(recs, TODAY)
    assert recs[3].date == date(2026, 9, 4)
    assert "медиана" in recs[3].date_res.note


def test_ambiguous_slash_without_any_hints_is_dmy():
    recs = _recs([("05/04/2026", "x")])
    cz.resolve_ambiguous_dates(recs, TODAY)
    assert recs[0].date == date(2026, 4, 5)


def test_ambiguous_slash_one_variant_in_future():
    recs = _recs([("09/11/2026", "x")])      # 9 ноября — будущее, 11 сентября — нет
    cz.resolve_ambiguous_dates(recs, TODAY)
    assert recs[0].date == date(2026, 9, 11)


def test_dot_date_is_never_flipped_because_of_future():
    recs = _recs([("10.11.2026", "x")])      # 10 ноября 2026 — будущее; 11 октября не подставляем
    cz.resolve_ambiguous_dates(recs, TODAY)
    assert recs[0].date is None and "в будущем" in recs[0].date_res.problem


def test_dot_date_plain_dmy_has_no_note():
    recs = _recs([("05.09.2026", "звонок"), ("17.09.2026", "звонок")])
    cz.resolve_ambiguous_dates(recs, TODAY)
    assert recs[0].date == date(2026, 9, 5) and recs[0].date_res.note is None


# ---------------------------------------------------------------- имя

@pytest.mark.parametrize("raw, expected", [
    ("сергей волков", "Сергей Волков"),
    ("ПАВЛОВА ирина", "Павлова Ирина"),
    ("Кузнецов  Андрей", "Кузнецов Андрей"),
    ("Смирнов Олег", "Смирнов Олег"),           # неразрывный пробел
    ("Смирнов Олег​", "Смирнов Олег"),          # zero-width space
    ("Пeтров Алексей", "Петров Алексей"),            # латинская e -> кириллическая
    ("влад", "Влад"),
    ("К. Белоусов", "К. Белоусов"),
    ("  анна-мария  петрова-водкина ", "Анна-Мария Петрова-Водкина"),
    ("Королёва Дарья", "Королёва Дарья"),            # ё сохраняем
    ("John SMITH", "John Smith"),                    # латиницу целиком не трогаем
    ("\tИван\tИванов\n", "Иван Иванов"),
])
def test_name_ok(raw, expected):
    assert cz.clean_name(raw) == expected


def test_name_homoglyph_fix_really_gives_cyrillic():
    assert cz.clean_name("Пeтров") == "Петров" and "e" not in cz.clean_name("Пeтров")


@pytest.mark.parametrize("raw", [None, "", "   ", "-", "123", "Нет", "не указано", "N/A", "​"])
def test_name_missing(raw):
    assert cz.clean_name(raw) is None


def test_names_compatible():
    assert cz.names_compatible("Сергей Волков", "Волков Сергей")
    assert cz.names_compatible("Королёва Дарья", "Королева Дарья")
    assert cz.names_compatible("К. Белоусов", "Белоусов Кирилл")
    assert not cz.names_compatible("Морозова Елена", "Белов Никита")
    assert cz.names_compatible(None, "Белов Никита")


def test_keeper_prefers_valid_then_earliest_then_fuller_name():
    a = cz.build_record(2, "К. Белоусов", "+79167703148", "27.09.2026", "звонок", TODAY)
    b = cz.build_record(3, "Белоусов Кирилл", "+79167703148", "27.09.2026", "сайт", TODAY)
    c = cz.build_record(4, "Белоусов Кирилл", "+79167703148", "20.09.2026", "email", TODAY)
    d = cz.build_record(5, None, "+79167703148", "01.09.2026", "email", TODAY)  # раньше всех, но без имени
    res = cz.process([a, b, c, d], TODAY)
    assert [r.row for r in res.clean] == [4]          # без проблем и самая ранняя
    assert {r.row for r in res.duplicates} == {2, 3, 5}
    res2 = cz.process([cz.build_record(2, "К. Белоусов", "+79167703148", "27.09.2026", "", TODAY),
                       cz.build_record(3, "Белоусов Кирилл", "+79167703148", "27.09.2026", "", TODAY)], TODAY)
    assert res2.clean[0].name == "Белоусов Кирилл"     # даты равны -> полное имя


# ---------------------------------------------------------------- исходный файл целиком

def read_xlsx(path):
    ws = openpyxl.load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


@pytest.fixture(scope="module")
def original_out(tmp_path_factory):
    out = tmp_path_factory.mktemp("orig")
    assert cz.main([str(ORIGINAL), "-o", str(out), "--today", TODAY.isoformat()]) == 0
    return {name: read_xlsx(out / f"{name}.xlsx") for name in ("clean", "problems", "duplicates")}


def test_original_counts(original_out):
    clean, problems, dups = original_out["clean"], original_out["problems"], original_out["duplicates"]
    assert (len(clean), len(dups), len(problems)) == (22, 14, 12)
    rows = [r["Строка"] for part in (clean, problems, dups) for r in part]
    assert sorted(rows) == list(range(2, 50))           # каждая строка ровно в одном файле


def test_original_clean_is_clean(original_out):
    import re
    clean = original_out["clean"]
    phones = [r["Телефон"] for r in clean]
    assert len(set(phones)) == len(phones)
    assert all(re.fullmatch(r"\+7\d{10}", p) for p in phones)
    assert all(re.fullmatch(r"2026-09-\d\d", r["Дата заявки"]) for r in clean)
    assert all(r["Имя"] == " ".join(w[:1].upper() + w[1:] for w in r["Имя"].split()) for r in clean)


def test_original_specific_decisions(original_out):
    clean = {r["Строка"]: r for r in original_out["clean"]}
    problems = {r["Строка"]: r for r in original_out["problems"]}
    dups = {r["Строка"]: r for r in original_out["duplicates"]}
    assert clean[2]["Дата заявки"] == "2026-09-03"      # 09/03 у «сайта» = ММ/ДД
    assert clean[8]["Дата заявки"] == "2026-09-07"      # 07/09 у telegram = ДД/ММ
    assert clean[35]["Дата заявки"] == "2026-09-04" and clean[36]["Дата заявки"] == "2026-09-04"
    assert clean[12]["Телефон"] == "+79174005219"       # номер из текста
    assert 48 in clean and dups[47]["Оставлена строка"] == 48   # «Белоусов Кирилл» полнее «К. Белоусов»
    assert dups[31]["Оставлена строка"] == 30           # Карпова: оставлена ранняя заявка
    assert "другой человек" in dups[9]["Примечание"]    # Морозова / Белов — один номер
    assert all("8-80X" in problems[r]["Причина"] for r in range(19, 24))  # 5 разных людей с 8-800
    assert "будущем" in problems[32]["Причина"]
    assert "нет имени" in problems[18]["Причина"]
    assert dups[16]["Оставлена строка"] == 15 and "нет года" in dups[16]["Примечание"]
    assert clean[26]["Имя"] == "Королёва Дарья" and clean[24]["Имя"] == "Петров Алексей"


# ---------------------------------------------------------------- сгенерированный файл

@pytest.fixture(scope="module")
def generated_out(tmp_path_factory):
    d = tmp_path_factory.mktemp("gen")
    src = gen.build(d / "zayavki_test2.xlsx")
    assert cz.main([str(src), "-o", str(d / "out"), "--today", gen.TODAY]) == 0
    return {name: {r["Строка"]: r for r in read_xlsx(d / "out" / f"{name}.xlsx")}
            for name in ("clean", "problems", "duplicates")}


@pytest.mark.parametrize("case", gen.CASES, ids=[c[0] for c in gen.CASES])
def test_generated_case(generated_out, case):
    key, *_raw, exp = case
    row = gen.row_of(key)
    clean, problems, dups = generated_out["clean"], generated_out["problems"], generated_out["duplicates"]
    where = [n for n, part in (("clean", clean), ("problem", problems), ("dup", dups)) if row in part]
    if exp[0] == "empty":
        assert where == []
        return
    assert where == [exp[0]], f"строка {row} ({key}) попала в {where}"
    if exp[0] == "clean":
        r = clean[row]
        assert (r["Имя"], r["Телефон"], r["Дата заявки"], r["Источник"]) == exp[1:]
    elif exp[0] == "problem":
        assert exp[1] in problems[row]["Причина"]
    else:
        assert dups[row]["Оставлена строка"] == gen.row_of(exp[1])


def test_generated_totals(generated_out):
    n = {k: len(v) for k, v in generated_out.items()}
    expected = {"clean": 0, "problem": 0, "dup": 0}
    for *_x, exp in gen.CASES:
        if exp[0] != "empty":
            expected[exp[0]] += 1
    assert (n["clean"], n["problems"], n["duplicates"]) == (expected["clean"], expected["problem"], expected["dup"])


def test_generated_conflict_flagged(generated_out):
    y1 = gen.row_of("Y1")
    assert "другое имя" in generated_out["clean"][y1]["Примечание"]
    assert generated_out["problems"][gen.row_of("D")]["Имя"] == "Козлов Олег"   # латинская K исправлена


# ---------------------------------------------------------------- CLI и устойчивость

def test_help_exits_zero(capsys):
    with pytest.raises(SystemExit) as e:
        cz.main(["--help"])
    assert e.value.code == 0 and "входной" in capsys.readouterr().out


def test_missing_file_and_bad_header(tmp_path, capsys):
    assert cz.main([str(tmp_path / "nope.xlsx"), "-o", str(tmp_path)]) == 2
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b", "c"])
    wb.active.append([1, 2, 3])
    wb.save(tmp_path / "bad.xlsx")
    assert cz.main([str(tmp_path / "bad.xlsx"), "-o", str(tmp_path)]) == 2
    assert "заголовков" in capsys.readouterr().err


def test_other_column_order_title_row_and_extra_column(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Заявки за август"])                           # строка-заголовок над таблицей
    ws.append(["Источник", " Дата заявки ", "Комментарий", "Телефон", "ФИО"])
    ws.append(["сайт", "01.08.2026", "перезвонить", "8 916 123 45 67", "иванов иван"])
    ws.append(["email", "02.08.2026", "", "+7 916 123-45-67", "Иван Иванов"])
    wb.save(tmp_path / "in.xlsx")
    assert cz.main([str(tmp_path / "in.xlsx"), "-o", str(tmp_path / "o"), "--today", "2026-10-11"]) == 0
    clean = read_xlsx(tmp_path / "o" / "clean.xlsx")
    assert len(clean) == 1
    assert (clean[0]["Имя"], clean[0]["Телефон"], clean[0]["Дата заявки"], clean[0]["Строка"]) == \
        ("Иванов Иван", "+79161234567", "2026-08-01", 3)


def test_idempotent_on_own_output(tmp_path):
    out1, out2 = tmp_path / "o1", tmp_path / "o2"
    assert cz.main([str(ORIGINAL), "-o", str(out1), "--today", "2026-10-11"]) == 0
    assert cz.main([str(out1 / "clean.xlsx"), "-o", str(out2), "--today", "2026-10-11"]) == 0
    first = [(r["Имя"], r["Телефон"], r["Дата заявки"], r["Источник"]) for r in read_xlsx(out1 / "clean.xlsx")]
    second = [(r["Имя"], r["Телефон"], r["Дата заявки"], r["Источник"]) for r in read_xlsx(out2 / "clean.xlsx")]
    assert first == second
    assert read_xlsx(out2 / "problems.xlsx") == [] and read_xlsx(out2 / "duplicates.xlsx") == []


def test_csv_written_with_bom_and_semicolon(tmp_path):
    assert cz.main([str(ORIGINAL), "-o", str(tmp_path), "--today", "2026-10-11"]) == 0
    raw = (tmp_path / "clean.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    first_line = raw.decode("utf-8-sig").splitlines()[0]
    assert first_line == "Имя;Телефон;Дата заявки;Источник;Строка;Примечание"
    assert len((tmp_path / "problems.csv").read_text(encoding="utf-8-sig").splitlines()) == 13
