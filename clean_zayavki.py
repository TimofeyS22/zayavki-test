#!/usr/bin/env python3
"""Чистка файла заявок (Имя, Телефон, Дата заявки, Источник).

Вход:  xlsx с колонками Имя / Телефон / Дата заявки / Источник
       (колонки ищутся по заголовку, порядок не важен).
Выход: clean.xlsx + clean.csv            — чистые уникальные записи;
       problems.xlsx + problems.csv      — проблемные строки с колонкой «Причина»;
       duplicates.xlsx + duplicates.csv  — удалённые дубли (что и с чем склеено).

Каждая входная строка попадает ровно в один из трёх файлов:
    входных строк = чистых + проблемных + удалённых дублей.

Правила нормализации и спорные решения описаны в README.md.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import median_low

try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Не установлен openpyxl: pip install -r requirements.txt")

NO_PHONE = "нет телефона"

# ---------------------------------------------------------------------------
# Общие утилиты
# ---------------------------------------------------------------------------

# Невидимые символы, которые \s не ловит: zero-width space/joiners, BOM, мягкий перенос.
_INVISIBLE = dict.fromkeys(map(ord, "​‌‍⁠﻿­"), None)


def clean_text(value) -> str:
    """Строка без невидимых символов, с одиночными обычными пробелами."""
    if value is None:
        return ""
    s = unicodedata.normalize("NFC", str(value)).translate(_INVISIBLE)
    return re.sub(r"\s+", " ", s).strip()  # \s в Python ловит и NBSP, и табы


def show(value) -> str:
    """Исходное значение ячейки в читаемом виде (для отчётов)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# ---------------------------------------------------------------------------
# Имя
# ---------------------------------------------------------------------------

# Латинские буквы, неотличимые на глаз от кириллических («Пeтров» с латинской e).
_HOMOGLYPHS = str.maketrans("aeopcxykAEOPCXYKBMHT", "аеорсхукАЕОРСХУКВМНТ")
_CYR = re.compile(r"[а-яё]", re.I)
_LAT = re.compile(r"[a-z]", re.I)
_LAT_HOMOGLYPH_ONLY = re.compile(r"^[^a-zA-Z]*(?:[aeopcxykAEOPCXYKBMHT][^a-zA-Z]*)+$")

NAME_PLACEHOLDERS = {
    "нет", "не указано", "не указан", "неизвестно", "не знаю", "без имени", "аноним",
    "имя", "клиент", "n/a", "na", "none", "null", "unknown", "noname",
}


def _fix_homoglyphs(word: str) -> str:
    # Меняем латиницу на кириллицу только в «смешанных» словах: есть кириллица,
    # а латинские буквы — только двойники кириллических. «John» не трогаем.
    if _CYR.search(word) and _LAT.search(word) and _LAT_HOMOGLYPH_ONLY.match(word):
        return word.translate(_HOMOGLYPHS)
    return word


def _cap(part: str) -> str:
    return part[:1].upper() + part[1:].lower()


def clean_name(raw) -> str | None:
    """Имя с заглавной буквы (каждое слово и часть через дефис), без лишних пробелов.
    None — если имени нет: пусто, нет ни одной буквы или заглушка вроде «нет»."""
    s = clean_text(raw).strip(" ,;:_-")
    if not s or not re.search(r"[^\W\d_]", s):
        return None
    if s.lower().strip(".") in NAME_PLACEHOLDERS:
        return None
    words = [_fix_homoglyphs(w) for w in s.split(" ")]
    return " ".join("-".join(_cap(p) for p in w.split("-")) for w in words)


def _name_tokens(name: str | None) -> list[str]:
    if not name:
        return []
    return [t for t in re.split(r"[\s.\-]+", name.lower().replace("ё", "е")) if t]


def names_compatible(a: str | None, b: str | None) -> bool:
    """Похожи ли имена на одного человека: те же слова в любом порядке, ё=е,
    инициал или сокращение совпадает с началом слова («К. Белоусов» ~ «Белоусов Кирилл»)."""
    ta, tb = _name_tokens(a), _name_tokens(b)
    if not ta or not tb:
        return True
    short, long_ = sorted((ta, tb), key=len)
    return all(any(u.startswith(t) or t.startswith(u) for u in long_) for t in short)


def name_completeness(name: str | None) -> int:
    """Сколько в имени полных слов (не инициалов)."""
    return sum(1 for t in (name or "").split() if len(t.strip(".")) > 1 and not t.endswith("."))


# ---------------------------------------------------------------------------
# Телефон
# ---------------------------------------------------------------------------

@dataclass
class PhoneResult:
    value: str | None          # +7XXXXXXXXXX или None
    problem: str | None = None  # причина для списка проблемных
    note: str | None = None     # пояснение для чистого файла


# Кусок текста, похожий на номер: цифры, между ними пробелы/дефисы/скобки/точки.
_PHONE_CHUNK = re.compile(r"\+?\d(?:[\s\-.()]*\d)*")
_EXT_WORDS = re.compile(r"доб|ext|вн\.|#", re.I)


def _national(digits: str, had_plus: bool) -> tuple[str | None, str | None]:
    """10 цифр национального номера или (None, причина)."""
    n = len(digits)
    if had_plus and not digits.startswith("7"):
        return None, "иностранный номер (код страны не +7)"
    if had_plus and n != 11:
        return None, f"не хватает цифр: {n}" if n < 11 else f"лишние цифры: {n}"
    if n == 11 and digits[0] in "78":
        return digits[1:], None
    if n == 10 and not had_plus:
        # Номер без кода страны: 9xx — мобильный, 3xx/4xx/80x–87x — коды регионов и служб.
        # «7…», «88…», «89…» из 10 цифр — это «7»/«8» + номер, в котором не хватает цифры
        # (кодов 88x/89x нет, а 7xx без +7 — это обрезанный номер).
        if digits[0] in "349" or (digits[0] == "8" and digits[1] not in "89"):
            return digits, None
        if digits[0] in "78":
            return None, f"не хватает цифры: 10 цифр, начинается с {digits[:2]}"
        return None, f"10 цифр, несуществующий код {digits[:3]}"
    if n < 10:
        return None, f"не хватает цифр: {n}"
    return None, f"лишние цифры: {n}"


def _check_national(nat: str) -> str | None:
    if nat[0] not in "34789":
        return f"несуществующий код +7 {nat[:3]}"
    if len(set(nat)) == 1:
        return "номер-заглушка из одинаковых цифр"
    if nat.startswith("80"):
        return "номер 8-80X — сервисная линия организации, не телефон клиента"
    return None


def normalize_phone(raw) -> PhoneResult:
    """Приводит телефон к +7XXXXXXXXXX. Если в ячейке не телефон клиента — value=None
    и причина (идёт в список проблемных как «нет телефона»)."""
    if raw is None or (isinstance(raw, str) and not clean_text(raw)):
        return PhoneResult(None, "пусто")
    if isinstance(raw, bool):
        return PhoneResult(None, f"не номер: «{raw}»")
    if isinstance(raw, (int, float)):  # Excel хранит номер числом: 89032218840
        if isinstance(raw, float) and not raw.is_integer():
            return PhoneResult(None, f"не номер: «{raw}»")
        text = str(int(raw))
    else:
        text = clean_text(raw)

    chunks = _PHONE_CHUNK.findall(text)
    if not chunks:
        return PhoneResult(None, f"не номер: «{text}»")

    first_problem = None
    for chunk in chunks:
        digits = re.sub(r"\D", "", chunk)
        if len(digits) < 5:  # «18:00», «доб. 45» и т.п. — не номер
            continue
        nat, problem = _national(digits, chunk.startswith("+"))
        if nat is not None:
            problem = _check_national(nat)
        if problem is None:
            notes = []
            if re.search(r"[^\W\d_]", text):
                notes.append(f"номер извлечён из текста «{text}»")
            if _EXT_WORDS.search(text):
                notes.append("добавочный номер отброшен")
            if len([c for c in chunks if len(re.sub(r"\D", "", c)) >= 10]) > 1:
                notes.append("в ячейке несколько номеров, взят первый корректный")
            return PhoneResult("+7" + nat, None, "; ".join(notes) or None)
        first_problem = first_problem or f"{problem}: «{text}»"

    return PhoneResult(None, first_problem or f"не номер: «{text}»")


# ---------------------------------------------------------------------------
# Дата
# ---------------------------------------------------------------------------

MIN_DATE = date(2000, 1, 1)
EXCEL_EPOCH = date(1899, 12, 30)  # система дат Excel 1900 (с учётом бага 29.02.1900)

_MONTHS = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5, "июн": 6, "июл": 7,
    "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7,
    "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}

_RE_YMD = re.compile(r"^(\d{4})[-./](\d{1,2})[-./](\d{1,2})$")
_RE_DMY = re.compile(r"^(\d{1,2})([-./])(\d{1,2})\2(\d{2}|\d{4})$")
_RE_TEXT_DMY = re.compile(r"^(\d{1,2})\s*([a-zа-яё]+)\.?,?\s*(\d{2}|\d{4})?$")
_RE_TEXT_MDY = re.compile(r"^([a-z]+)\.?\s+(\d{1,2}),?\s*(\d{4})?$")
_RE_TIME_TAIL = re.compile(r"(?:[ t]+\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)$")
_RE_YEAR_WORD = re.compile(r"(?<=\d)\s*(?:года|год|г)\.?$")  # «2026 г.», но не «авг»


@dataclass
class DateResult:
    value: date | None
    problem: str | None = None
    note: str | None = None
    # Для «09/03/2026»: оба варианта (ДД/ММ, ММ/ДД) возможны — решаем по всему файлу.
    ambiguous: tuple[date, date] | None = None
    sep: str | None = None              # разделитель Д/М-даты: "/", ".", "-"
    order_evidence: str | None = None   # "DMY"/"MDY", если порядок однозначен (17/09, 09/16)


def _year(y: str) -> int:
    return 2000 + int(y) if len(y) == 2 else int(y)  # 26 -> 2026


def _make(y: int, m: int, d: int) -> date | None:
    try:
        return date(y, m, d)
    except ValueError:
        return None


def check_range(d: date, today: date) -> str | None:
    if d > today:
        return f"дата в будущем: {d.isoformat()}"
    if d < MIN_DATE:
        return f"неправдоподобно старая дата: {d.isoformat()}"
    return None


def parse_date(raw, today: date) -> DateResult:
    """Разбирает дату заявки. Неоднозначные Д/М-даты возвращаются с ambiguous=(ДМГ, МДГ),
    окончательно их решает resolve_ambiguous_dates() по всему файлу."""
    if raw is None or (isinstance(raw, str) and not clean_text(raw)):
        return DateResult(None, "пусто")

    if isinstance(raw, datetime):
        d = raw.date()
    elif isinstance(raw, date):
        d = raw
    elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
        d = _from_excel_serial(raw)
        if d is None:
            return DateResult(None, f"не распознана: число «{show(raw)}»")
    else:
        return _parse_date_text(clean_text(raw), today)

    problem = check_range(d, today)
    return DateResult(None if problem else d, problem)


def _from_excel_serial(n) -> date | None:
    # 36526 = 2000-01-01, 73051 = 2100-01-01. Вне диапазона — это не дата.
    if 36526 <= n < 73051:
        return EXCEL_EPOCH + timedelta(days=int(n))
    return None


def _parse_date_text(text: str, today: date) -> DateResult:
    s = text.lower().replace("ё", "е")
    s = _RE_TIME_TAIL.sub("", s).strip()
    s = _RE_YEAR_WORD.sub("", s).strip()

    if re.fullmatch(r"\d{5}(?:\.0)?", s):  # Excel-дата числом, сохранённая как текст
        d = _from_excel_serial(float(s))
        if d:
            return _finish(d, today)

    m = _RE_YMD.match(s)
    if m:
        d = _make(int(m[1]), int(m[2]), int(m[3]))
        return _finish(d, today, text)

    m = _RE_DMY.match(s)
    if m:
        a, sep, b, y = int(m[1]), m[2], int(m[3]), _year(m[4])
        if a > 12 or b > 12 or a == b:
            # Порядок однозначен: 17/09 — день первый, 09/16 — месяц первый.
            order = "MDY" if b > 12 else "DMY"
            # _make(год, МЕСЯЦ, ДЕНЬ): для ДД/ММ месяц = b, для ММ/ДД месяц = a
            d = _make(y, a, b) if order == "MDY" else _make(y, b, a)
            res = _finish(d, today, text)
            if a != b:
                res.sep, res.order_evidence = sep, order
            if order == "MDY" and res.value:
                res.note = f"дата «{text}» записана как ММ{sep}ДД"
            return res
        dmy, mdy = _make(y, b, a), _make(y, a, b)
        return DateResult(None, ambiguous=(dmy, mdy), sep=sep)

    for rx, day_i, mon_i in ((_RE_TEXT_DMY, 1, 2), (_RE_TEXT_MDY, 2, 1)):
        m = rx.match(s)
        if m and m[mon_i][:3] in _MONTHS:
            if not m[3]:
                return DateResult(None, f"нет года: «{text}»")
            d = _make(_year(m[3]), _MONTHS[m[mon_i][:3]], int(m[day_i]))
            return _finish(d, today, text)

    return DateResult(None, f"не распознана: «{text}»")


def _finish(d: date | None, today: date, text: str = "") -> DateResult:
    if d is None:
        return DateResult(None, f"несуществующая дата: «{text}»")
    problem = check_range(d, today)
    return DateResult(None if problem else d, problem)


# ---------------------------------------------------------------------------
# Записи, разрешение неоднозначных дат, дедупликация
# ---------------------------------------------------------------------------

@dataclass
class Record:
    row: int                       # номер строки в исходном Excel (заголовок = 1)
    raw_name: object
    raw_phone: object
    raw_date: object
    raw_source: object
    name: str | None = None
    phone: PhoneResult = None
    date_res: DateResult = None
    source: str = ""
    problems: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    kept_row: int | None = None    # для удалённого дубля — какая строка оставлена

    @property
    def date(self) -> date | None:
        return self.date_res.value


def build_record(row: int, name, phone, dt, source, today: date) -> Record:
    rec = Record(row, name, phone, dt, source)
    rec.name = clean_name(name)
    rec.phone = normalize_phone(phone)
    rec.date_res = parse_date(dt, today)
    rec.source = clean_text(source).lower()
    return rec


def resolve_ambiguous_dates(records: list[Record], today: date) -> None:
    """Даты вида «09/03/2026», где и день, и месяц <= 12: 9 марта или 3 сентября?

    Сигнал «источника»: если у этого источника ВСЕ однозначные даты с тем же
    разделителем записаны одним порядком (сайт: 09/16, 09/24 -> ММ/ДД; telegram:
    17/09, 22/09 -> ДД/ММ), неоднозначные даты этого источника читаем так же.

    «.» и «-» (российская запись): ДД.ММ, если только сигнал источника не говорит
    обратное. Дату не «переворачиваем» из-за того, что она в будущем — это проблема.

    «/» (может быть и американской записью ММ/ДД), по шагам:
      1) сигнал источника;
      2) если один из вариантов невозможен (в будущем / раньше 2000 г.) — другой;
      3) вариант, ближайший к медиане остальных дат файла;
      4) ДД/ММ (российская норма)."""
    evidence: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for r in records:
        if r.date_res.order_evidence:
            evidence[(r.source, r.date_res.sep)][r.date_res.order_evidence] += 1
    known = sorted(r.date for r in records if r.date)
    ref = median_low(known) if known else None

    for r in records:
        dr = r.date_res
        if not dr.ambiguous:
            continue
        raw = clean_text(r.raw_date)
        cand = dict(zip(("DMY", "MDY"), dr.ambiguous))
        votes = evidence[(r.source, dr.sep)]
        by_source = None
        if votes["DMY"] and not votes["MDY"]:
            by_source = "DMY"
        elif votes["MDY"] and not votes["DMY"]:
            by_source = "MDY"
        how = None
        if by_source:
            choice = by_source
            how = f"все однозначные даты источника «{r.source or 'без источника'}» с «{dr.sep}» записаны так"
            if dr.sep != "/" and choice == "DMY":
                how = None  # обычная российская запись, пояснять нечего
        elif dr.sep != "/":
            choice = "DMY"
        else:
            feasible = [k for k in cand if not check_range(cand[k], today)]
            if len(feasible) == 1:
                choice, how = feasible[0], "второй вариант невозможен (в будущем или слишком старый)"
            elif feasible and ref is not None:
                choice = min(feasible, key=lambda k: (abs((cand[k] - ref).days), k != "DMY"))
                how = f"этот вариант ближе к остальным датам файла (медиана {ref.isoformat()})"
            else:
                choice, how = "DMY", "подсказок нет, по умолчанию российский порядок"
        problem = check_range(cand[choice], today)
        if problem:
            dr.problem = problem
            continue
        dr.value = cand[choice]
        if how:
            fmt = f"ДД{dr.sep}ММ" if choice == "DMY" else f"ММ{dr.sep}ДД"
            dr.note = f"дата «{raw}» неоднозначна, принят формат {fmt}: {how}"


def _collect_problems(rec: Record) -> None:
    if rec.phone.value is None:
        rec.problems.append(f"{NO_PHONE} ({rec.phone.problem})")
    if rec.date is None:
        rec.problems.append(f"битая дата ({rec.date_res.problem})")
    if rec.name is None:
        rec.problems.append("нет имени")
    for note in (rec.phone.note, rec.date_res.note):
        if note:
            rec.notes.append(note)


def keeper_rank(rec: Record):
    """Какую запись оставить из группы с одним телефоном (меньше = лучше):
    1) запись без проблем; 2) самая ранняя дата заявки (первое обращение:
    дата и источник первого касания); 3) более полное имя; 4) раньше в файле."""
    return (bool(rec.problems), rec.date or date.max, -name_completeness(rec.name), rec.row)


@dataclass
class Result:
    clean: list[Record]
    problems: list[Record]
    duplicates: list[Record]
    total: int
    skipped_empty: int = 0


def process(records: list[Record], today: date, skipped_empty: int = 0) -> Result:
    resolve_ambiguous_dates(records, today)
    for rec in records:
        _collect_problems(rec)

    groups: dict[str, list[Record]] = defaultdict(list)
    for rec in records:
        if rec.phone.value:
            groups[rec.phone.value].append(rec)

    keepers: list[Record] = [r for r in records if not r.phone.value]  # без телефона не с чем склеивать
    duplicates: list[Record] = []
    for group in groups.values():
        group.sort(key=keeper_rank)
        keeper, rest = group[0], group[1:]
        keepers.append(keeper)
        refs = []
        for dup in rest:
            dup.kept_row = keeper.row
            same = names_compatible(keeper.name, dup.name)
            if not same:
                dup.notes.append(
                    f"имя отличается от оставленной записи («{keeper.name}») — возможно, другой человек")
            head = [f"дубль по телефону {keeper.phone.value}, оставлена строка {keeper.row}"]
            if dup.problems:
                head.append("у самой строки: " + "; ".join(dup.problems))
            dup.notes[:0] = head
            label = f"стр. {dup.row}" + ("" if same else f" «{dup.name or 'без имени'}» — другое имя, проверьте")
            refs.append(label)
            duplicates.append(dup)
        if refs:
            keeper.notes.insert(0, "дубли по телефону: " + ", ".join(refs))

    keepers.sort(key=lambda r: r.row)
    duplicates.sort(key=lambda r: r.row)
    return Result(
        clean=[r for r in keepers if not r.problems],
        problems=[r for r in keepers if r.problems],
        duplicates=duplicates,
        total=len(records),
        skipped_empty=skipped_empty,
    )


# ---------------------------------------------------------------------------
# Чтение и запись файлов
# ---------------------------------------------------------------------------

HEADER_KEYS = {
    "name": ("имя", "фио", "name", "клиент"),
    "phone": ("телефон", "тел", "phone", "моб"),
    "date": ("дата", "date"),
    "source": ("источник", "source", "канал"),
}


class InputError(Exception):
    pass


def _match_header(cell) -> str | None:
    h = clean_text(cell).lower()
    if not h:
        return None
    for key, aliases in HEADER_KEYS.items():
        if any(h.startswith(a) for a in aliases):
            return key
    return None


def read_rows(path: Path, sheet: str | None = None):
    """Возвращает (имя листа, [(номер строки, имя, телефон, дата, источник)], пустых строк)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except FileNotFoundError:
        raise InputError(f"файл не найден: {path}")
    except Exception as e:  # битый/не-xlsx файл
        raise InputError(f"не удалось открыть {path} как .xlsx: {e}")
    if sheet:
        if sheet not in wb.sheetnames:
            raise InputError(f"нет листа «{sheet}», есть: {', '.join(wb.sheetnames)}")
        ws = wb[sheet]
    else:
        ws = wb.worksheets[0]

    rows = list(ws.iter_rows(values_only=True))
    header_idx, cols = None, {}
    for i, row in enumerate(rows[:20]):
        found = {}
        for j, cell in enumerate(row):
            key = _match_header(cell)
            if key and key not in found:
                found[key] = j
        if {"name", "phone", "date"} <= found.keys():
            header_idx, cols = i, found
            break
    if header_idx is None:
        raise InputError("не найдена строка заголовков с колонками Имя, Телефон, Дата заявки")

    def get(row, key):
        j = cols.get(key)
        return row[j] if j is not None and j < len(row) else None

    data, empty = [], 0
    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        vals = [get(row, k) for k in ("name", "phone", "date", "source")]
        if all(v is None or (isinstance(v, str) and not clean_text(v)) for v in vals):
            empty += 1
            continue
        data.append((i, *vals))
    return ws.title, data, empty


def _fmt_date(rec: Record) -> str:
    return rec.date.isoformat() if rec.date else ""


CLEAN_COLS = ["Имя", "Телефон", "Дата заявки", "Источник", "Строка", "Примечание"]
PROBLEM_COLS = ["Строка", "Имя", "Телефон", "Дата заявки", "Источник", "Причина",
                "Имя (исходное)", "Телефон (исходный)", "Дата (исходная)", "Примечание"]
DUP_COLS = ["Строка", "Имя", "Телефон", "Дата заявки", "Источник", "Оставлена строка",
            "Имя (исходное)", "Телефон (исходный)", "Дата (исходная)", "Примечание"]


def clean_row(r: Record) -> list:
    return [r.name, r.phone.value, _fmt_date(r), r.source, r.row, "; ".join(r.notes)]


def problem_row(r: Record) -> list:
    return [r.row, r.name or "", r.phone.value or NO_PHONE, _fmt_date(r), r.source,
            "; ".join(r.problems), show(r.raw_name), show(r.raw_phone), show(r.raw_date),
            "; ".join(r.notes)]


def dup_row(r: Record) -> list:
    return [r.row, r.name or "", r.phone.value, _fmt_date(r), r.source, r.kept_row,
            show(r.raw_name), show(r.raw_phone), show(r.raw_date), "; ".join(r.notes)]


def write_xlsx(path: Path, title: str, header: list[str], rows: list[list]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(header)
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(header, start=1):
        width = max([len(str(col))] + [len(str(r[j - 1] or "")) for r in rows])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 8), 70)
    # Телефоны и даты пишем текстом, чтобы Excel не превращал их в числа/свои форматы.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.number_format = "@"
    wb.save(path)


def write_csv(path: Path, header: list[str], rows: list[list], sep: str) -> None:
    # utf-8-sig: Excel корректно откроет кириллицу двойным кликом.
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=sep)
        w.writerow(header)
        w.writerows(rows)


def write_outputs(result: Result, clean_path: Path, problems_path: Path, dups_path: Path,
                  sep: str) -> None:
    for path, title, header, rows in (
        (clean_path, "Чистые", CLEAN_COLS, [clean_row(r) for r in result.clean]),
        (problems_path, "Проблемные", PROBLEM_COLS, [problem_row(r) for r in result.problems]),
        (dups_path, "Дубли", DUP_COLS, [dup_row(r) for r in result.duplicates]),
    ):
        path.parent.mkdir(parents=True, exist_ok=True)
        write_xlsx(path, title, header, rows)
        write_csv(path.with_suffix(".csv"), header, rows, sep)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def run(input_path: Path, today: date, sheet: str | None = None) -> tuple[str, Result]:
    sheet_name, data, empty = read_rows(input_path, sheet)
    records = [build_record(row, n, p, d, s, today) for row, n, p, d, s in data]
    return sheet_name, process(records, today, empty)


def summary(result: Result) -> str:
    cats = Counter()
    for r in result.problems:
        for p in r.problems:
            cats[p.split(" (")[0]] += 1
    n_clean, n_dup, n_prob = len(result.clean), len(result.duplicates), len(result.problems)
    ok = n_clean + n_dup + n_prob == result.total
    lines = [
        f"Строк с данными на входе: {result.total}"
        + (f" (пустых строк пропущено: {result.skipped_empty})" if result.skipped_empty else ""),
        f"  чистых уникальных записей: {n_clean}",
        f"  удалено дублей по телефону: {n_dup}",
        f"  проблемных строк:           {n_prob}"
        + (" — " + ", ".join(f"{k}: {v}" for k, v in cats.most_common()) if cats else ""),
        f"Сверка: {n_clean} + {n_dup} + {n_prob} = {n_clean + n_dup + n_prob} "
        + ("== вход, OK" if ok else f"!= {result.total} — ОШИБКА"),
    ]
    return "\n".join(lines)


def parse_args(argv=None) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        prog="clean_zayavki.py",
        description="Чистка файла заявок: телефоны -> +7XXXXXXXXXX, даты -> ГГГГ-ММ-ДД, "
                    "имена с заглавной, дубли по телефону убираются, проблемные строки — "
                    "в отдельный файл с причиной.",
        epilog="Пример: python clean_zayavki.py zayavki.xlsx -o output",
    )
    p.add_argument("input", type=Path, help="входной .xlsx (колонки Имя, Телефон, Дата заявки, Источник)")
    p.add_argument("-o", "--out-dir", type=Path, default=None,
                   help="папка для результатов (по умолчанию: output рядом со входным файлом)")
    p.add_argument("--clean", type=Path, help="путь к чистому файлу (по умолчанию <out-dir>/clean.xlsx)")
    p.add_argument("--problems", type=Path, help="путь к файлу проблемных (по умолчанию <out-dir>/problems.xlsx)")
    p.add_argument("--duplicates", type=Path,
                   help="путь к отчёту об удалённых дублях (по умолчанию <out-dir>/duplicates.xlsx)")
    p.add_argument("--sheet", help="имя листа (по умолчанию первый)")
    p.add_argument("--today", type=date.fromisoformat, default=None,
                   help="«сегодня» для проверки дат в будущем, ГГГГ-ММ-ДД (по умолчанию текущая дата)")
    p.add_argument("--csv-sep", default=";", help="разделитель CSV (по умолчанию «;» — для русского Excel)")
    return p.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    today = args.today or date.today()
    out_dir = args.out_dir or args.input.parent / "output"
    clean_path = args.clean or out_dir / "clean.xlsx"
    problems_path = args.problems or out_dir / "problems.xlsx"
    dups_path = args.duplicates or out_dir / "duplicates.xlsx"
    for pth in (clean_path, problems_path, dups_path):
        if pth.suffix.lower() != ".xlsx":
            print(f"Ошибка: выходной файл должен быть .xlsx: {pth}", file=sys.stderr)
            return 2
    try:
        sheet_name, result = run(args.input, today, args.sheet)
    except InputError as e:
        print(f"Ошибка: {e}", file=sys.stderr)
        return 2
    write_outputs(result, clean_path, problems_path, dups_path, args.csv_sep)
    print(f"Вход: {args.input} (лист «{sheet_name}»), сегодня = {today.isoformat()}")
    print(summary(result))
    print("Файлы:")
    for pth in (clean_path, problems_path, dups_path):
        print(f"  {pth}  (+ {pth.with_suffix('.csv').name})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
